"""
Offer Generator - The orchestrator that brings everything together.

Philosophy: This is where data transforms into beautiful presentations.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pathlib import Path
from datetime import date, timedelta
from dateutil.relativedelta import relativedelta

from src.models import Offer
from src.calculator import PriceCalculator
from src.charts import ChartGenerator, TableGenerator


class OfferGenerator:
    """Generates PowerPoint offers from data"""

    def __init__(self, offer: Offer, template_path: str):
        self.offer = offer
        self.template_path = template_path
        self.calculator = PriceCalculator(offer)
        self.chart_gen = ChartGenerator(offer)
        self.table_gen = TableGenerator(offer)

    def generate(self, output_path: str) -> str:
        """
        Generate the offer presentation.

        Returns:
            Path to the generated file
        """
        print(f"🎨 Generating offer for {self.offer.customer.name}...")

        # Load template
        prs = Presentation(self.template_path)

        # Calculate pricing
        breakdown = self.calculator.calculate_breakdown()

        # Replace placeholders in each slide
        self._fill_cover_slide(prs.slides[0], breakdown)
        self._fill_executive_summary(prs.slides[1], breakdown)
        self._fill_team_composition(prs.slides[2])
        self._fill_timeline(prs.slides[3])
        self._fill_budget_breakdown(prs.slides[4])
        self._fill_pricing_details(prs.slides[5], breakdown)
        self._fill_next_steps(prs.slides[6])

        # Save
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path)

        print(f"✨ Offer generated: {output_path}")
        return output_path

    def _replace_text_in_shape(self, shape, replacements: dict):
        """Replace placeholder text in a shape"""
        if not shape.has_text_frame:
            return

        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                for placeholder, value in replacements.items():
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, value)

    def _replace_text_in_slide(self, slide, replacements: dict):
        """Replace all placeholders in a slide"""
        for shape in slide.shapes:
            self._replace_text_in_shape(shape, replacements)

    def _fill_cover_slide(self, slide, breakdown):
        """Fill cover slide with data"""
        print("  ✓ Cover slide")

        replacements = {
            '{{project_name}}': self.offer.project.name,
            '{{customer_name}}': self.offer.customer.name,
            '{{offer_date}}': self.offer.offer_date.strftime('%d.%m.%Y')
        }

        self._replace_text_in_slide(slide, replacements)

    def _fill_executive_summary(self, slide, breakdown):
        """Fill executive summary slide"""
        print("  ✓ Executive summary")

        # Calculate end date
        end_date = self.offer.project.start_date + relativedelta(
            months=self.offer.project.duration_months
        )

        # Project description
        description = self.offer.project.description or f"Softwareentwicklungsprojekt für {self.offer.customer.name}"

        replacements = {
            '{{project_description}}': description,
            '{{duration_months}}': str(self.offer.project.duration_months),
            '{{start_date}}': self.offer.project.start_date.strftime('%d.%m.%Y'),
            '{{team_size}}': str(self.calculator.get_team_size()),
            '{{total_price}}': breakdown.format_currency(breakdown.total)
        }

        self._replace_text_in_slide(slide, replacements)

    def _fill_team_composition(self, slide):
        """Fill team composition slide with table"""
        print("  ✓ Team composition")

        # Remove placeholder text
        for shape in slide.shapes:
            if shape.has_text_frame and '{{TEAM_TABLE}}' in shape.text_frame.text:
                # Get position and size
                left = shape.left
                top = shape.top
                width = shape.width
                height = shape.height

                # Remove placeholder
                sp = shape.element
                sp.getparent().remove(sp)

                # Add table
                self.table_gen.create_team_table(slide, left, top, width, height)
                break

    def _fill_timeline(self, slide):
        """Fill timeline slide with chart"""
        print("  ✓ Timeline")

        # Remove placeholder and add chart
        for shape in slide.shapes:
            if shape.has_text_frame and '{{TIMELINE_CHART}}' in shape.text_frame.text:
                left = shape.left
                top = shape.top
                width = shape.width
                height = shape.height

                sp = shape.element
                sp.getparent().remove(sp)

                self.chart_gen.create_timeline_chart(slide, left, top, width, height)
                break

    def _fill_budget_breakdown(self, slide):
        """Fill budget breakdown slide with chart"""
        print("  ✓ Budget breakdown")

        for shape in slide.shapes:
            if shape.has_text_frame and '{{BUDGET_CHART}}' in shape.text_frame.text:
                left = shape.left
                top = shape.top
                width = shape.width
                height = shape.height

                sp = shape.element
                sp.getparent().remove(sp)

                self.chart_gen.create_budget_pie_chart(slide, left, top, width, height)
                break

    def _fill_pricing_details(self, slide, breakdown):
        """Fill pricing details slide with table"""
        print("  ✓ Pricing details")

        for shape in slide.shapes:
            if shape.has_text_frame and '{{PRICING_TABLE}}' in shape.text_frame.text:
                left = shape.left
                top = shape.top
                width = shape.width
                height = shape.height

                sp = shape.element
                sp.getparent().remove(sp)

                self.table_gen.create_pricing_table(slide, left, top, width, height)
                break

    def _fill_next_steps(self, slide):
        """Fill next steps slide"""
        print("  ✓ Next steps")

        replacements = {
            '{{start_date}}': self.offer.project.start_date.strftime('%d.%m.%Y')
        }

        self._replace_text_in_slide(slide, replacements)


def generate_offer(offer: Offer, template_path: str, output_path: str) -> str:
    """
    Convenience function to generate an offer.

    Args:
        offer: The offer data
        template_path: Path to the PowerPoint template
        output_path: Path where the generated offer should be saved

    Returns:
        Path to the generated file
    """
    generator = OfferGenerator(offer, template_path)
    return generator.generate(output_path)


# Example usage
if __name__ == "__main__":
    import yaml
    import sys
    from src.models import Offer

    # Load example data
    with open('data/example.yaml', 'r') as f:
        data = yaml.safe_load(f)

    offer = Offer(**data)

    # Generate offer
    output_path = f"output/{offer.customer.name.lower().replace(' ', '-')}-offer-{date.today().isoformat()}.pptx"

    generate_offer(
        offer=offer,
        template_path="templates/offer_template.pptx",
        output_path=output_path
    )

    print(f"\n🎉 Success! Open the file:")
    print(f"   {output_path}")
