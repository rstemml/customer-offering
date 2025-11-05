"""
Template Builder - Creates the initial PowerPoint template.

Philosophy: Templates should be beautiful, structured, and placeholder-aware.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path


class TemplateBuilder:
    """Builds a professional PowerPoint template"""

    # Design System - Consistent colors and spacing
    COLORS = {
        'primary': RGBColor(0, 102, 204),      # Professional Blue
        'secondary': RGBColor(102, 102, 102),  # Dark Gray
        'accent': RGBColor(255, 153, 0),       # Orange
        'light_gray': RGBColor(242, 242, 242), # Light Gray Background
        'white': RGBColor(255, 255, 255),
        'black': RGBColor(0, 0, 0)
    }

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(10)
        self.prs.slide_height = Inches(7.5)

    def _add_blank_slide(self):
        """Add a blank slide"""
        blank_layout = self.prs.slide_layouts[6]  # Blank layout
        return self.prs.slides.add_slide(blank_layout)

    def _add_header(self, slide, title_text):
        """Add consistent header to slide"""
        # Header background
        header = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0),
            Inches(10), Inches(1)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = self.COLORS['primary']
        header.line.fill.background()

        # Title text
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.25),
            Inches(9), Inches(0.5)
        )
        title_frame = title_box.text_frame
        title_frame.text = title_text
        title_p = title_frame.paragraphs[0]
        title_p.font.size = Pt(32)
        title_p.font.bold = True
        title_p.font.color.rgb = self.COLORS['white']

        return slide

    def build_cover_slide(self):
        """Slide 1: Cover / Deckblatt"""
        slide = self._add_blank_slide()

        # Background color
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = self.COLORS['primary']

        # Company logo placeholder (optional)
        logo_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5),
            Inches(2), Inches(0.5)
        )
        logo_frame = logo_box.text_frame
        logo_frame.text = "[YOUR LOGO]"
        logo_p = logo_frame.paragraphs[0]
        logo_p.font.size = Pt(14)
        logo_p.font.color.rgb = self.COLORS['white']

        # Main title
        title_box = slide.shapes.add_textbox(
            Inches(1), Inches(2.5),
            Inches(8), Inches(1.5)
        )
        title_frame = title_box.text_frame
        title_frame.text = "{{project_name}}"
        title_frame.word_wrap = True
        title_p = title_frame.paragraphs[0]
        title_p.font.size = Pt(54)
        title_p.font.bold = True
        title_p.font.color.rgb = self.COLORS['white']
        title_p.alignment = PP_ALIGN.CENTER

        # Subtitle
        subtitle_box = slide.shapes.add_textbox(
            Inches(1), Inches(4.2),
            Inches(8), Inches(0.8)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = "Angebot für {{customer_name}}"
        subtitle_p = subtitle_frame.paragraphs[0]
        subtitle_p.font.size = Pt(28)
        subtitle_p.font.color.rgb = self.COLORS['white']
        subtitle_p.alignment = PP_ALIGN.CENTER

        # Date
        date_box = slide.shapes.add_textbox(
            Inches(1), Inches(6.5),
            Inches(8), Inches(0.5)
        )
        date_frame = date_box.text_frame
        date_frame.text = "{{offer_date}}"
        date_p = date_frame.paragraphs[0]
        date_p.font.size = Pt(18)
        date_p.font.color.rgb = self.COLORS['white']
        date_p.alignment = PP_ALIGN.CENTER

        return slide

    def build_executive_summary(self):
        """Slide 2: Executive Summary"""
        slide = self._add_blank_slide()
        self._add_header(slide, "Executive Summary")

        # Content box
        content_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.5),
            Inches(9), Inches(5.5)
        )
        tf = content_box.text_frame
        tf.word_wrap = True

        # Project Overview
        p1 = tf.paragraphs[0]
        p1.text = "Projekt"
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = self.COLORS['primary']
        p1.space_after = Pt(10)

        p2 = tf.add_paragraph()
        p2.text = "{{project_description}}"
        p2.font.size = Pt(16)
        p2.space_after = Pt(20)

        # Key Facts
        p3 = tf.add_paragraph()
        p3.text = "Key Facts"
        p3.font.size = Pt(24)
        p3.font.bold = True
        p3.font.color.rgb = self.COLORS['primary']
        p3.space_after = Pt(10)

        facts = [
            "Dauer: {{duration_months}} Monate",
            "Start: {{start_date}}",
            "Team: {{team_size}} Personen",
            "Investition: {{total_price}} EUR"
        ]

        for fact in facts:
            p = tf.add_paragraph()
            p.text = f"• {fact}"
            p.font.size = Pt(16)
            p.level = 0

        return slide

    def build_team_composition(self):
        """Slide 3: Team Composition"""
        slide = self._add_blank_slide()
        self._add_header(slide, "Team Zusammensetzung")

        # Table will be inserted here by the generator
        # Placeholder for positioning
        table_placeholder = slide.shapes.add_textbox(
            Inches(1), Inches(2),
            Inches(8), Inches(4)
        )
        tf = table_placeholder.text_frame
        tf.text = "{{TEAM_TABLE}}"

        return slide

    def build_timeline(self):
        """Slide 4: Project Timeline"""
        slide = self._add_blank_slide()
        self._add_header(slide, "Projekt Zeitplan")

        # Chart placeholder
        chart_placeholder = slide.shapes.add_textbox(
            Inches(1), Inches(2),
            Inches(8), Inches(4.5)
        )
        tf = chart_placeholder.text_frame
        tf.text = "{{TIMELINE_CHART}}"

        return slide

    def build_budget_breakdown(self):
        """Slide 5: Budget Breakdown"""
        slide = self._add_blank_slide()
        self._add_header(slide, "Budget Verteilung")

        # Chart placeholder
        chart_placeholder = slide.shapes.add_textbox(
            Inches(1.5), Inches(2),
            Inches(7), Inches(4.5)
        )
        tf = chart_placeholder.text_frame
        tf.text = "{{BUDGET_CHART}}"

        return slide

    def build_pricing_details(self):
        """Slide 6: Pricing Details"""
        slide = self._add_blank_slide()
        self._add_header(slide, "Preiskalkulation")

        # Pricing table placeholder
        table_placeholder = slide.shapes.add_textbox(
            Inches(2), Inches(2),
            Inches(6), Inches(4.5)
        )
        tf = table_placeholder.text_frame
        tf.text = "{{PRICING_TABLE}}"

        return slide

    def build_next_steps(self):
        """Slide 7: Next Steps"""
        slide = self._add_blank_slide()
        self._add_header(slide, "Nächste Schritte")

        # Content
        content_box = slide.shapes.add_textbox(
            Inches(1.5), Inches(2),
            Inches(7), Inches(4)
        )
        tf = content_box.text_frame
        tf.word_wrap = True

        steps = [
            "Angebot prüfen und Feedback geben",
            "Kick-off Meeting vereinbaren",
            "Vertrag unterschreiben",
            "Projekt Start: {{start_date}}"
        ]

        for i, step in enumerate(steps, 1):
            p = tf.add_paragraph() if i > 1 else tf.paragraphs[0]
            p.text = f"{i}. {step}"
            p.font.size = Pt(22)
            p.space_after = Pt(20)

        # Contact information
        contact_box = slide.shapes.add_textbox(
            Inches(1.5), Inches(5.5),
            Inches(7), Inches(1.5)
        )
        tf_contact = contact_box.text_frame

        p1 = tf_contact.paragraphs[0]
        p1.text = "Kontakt"
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = self.COLORS['primary']

        p2 = tf_contact.add_paragraph()
        p2.text = "Bei Fragen stehen wir Ihnen gerne zur Verfügung."
        p2.font.size = Pt(14)

        return slide

    def build_template(self, output_path: str):
        """Build complete template"""
        print("🎨 Building template...")

        self.build_cover_slide()
        print("  ✓ Cover slide")

        self.build_executive_summary()
        print("  ✓ Executive summary")

        self.build_team_composition()
        print("  ✓ Team composition")

        self.build_timeline()
        print("  ✓ Timeline")

        self.build_budget_breakdown()
        print("  ✓ Budget breakdown")

        self.build_pricing_details()
        print("  ✓ Pricing details")

        self.build_next_steps()
        print("  ✓ Next steps")

        # Save
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(output_path)
        print(f"\n✨ Template created: {output_path}")


def create_default_template():
    """Create the default template"""
    builder = TemplateBuilder()
    builder.build_template("templates/offer_template.pptx")


if __name__ == "__main__":
    create_default_template()
